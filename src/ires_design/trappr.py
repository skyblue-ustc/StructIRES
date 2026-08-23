"""Normalize the IRES-TrAPPr supplementary workbooks without erasing assay context.

The supplementary tables mix binary activity calls, continuous translation
efficiency measurements, mutational scans, stress conditions, and validation
summaries.  This module deliberately keeps one row per measurement and builds a
second sequence-level table only as a derived, provenance-preserving view.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path
from statistics import fmean
from typing import Any, Iterable

from .schemas import normalize_rna


ASSAY_CONTEXT = "direct RNA A-cap/G-cap nascent-translation MPRA"
MEASUREMENT_FIELDS = (
    "record_id",
    "source_table",
    "source_row",
    "construct_name",
    "sequence",
    "sequence_sha256",
    "length",
    "alphabet_valid",
    "assay_context",
    "condition",
    "group",
    "ires_type",
    "organism",
    "activity_label",
    "activity_status",
    "te_ires",
    "te_cdi",
    "fold_difference_from_wt",
    "position",
    "is_control",
    "notes",
    "pmid",
)


@dataclass(frozen=True)
class TrapprMeasurement:
    record_id: str
    source_table: str
    source_row: int
    construct_name: str
    sequence: str
    sequence_sha256: str
    length: int
    alphabet_valid: bool
    assay_context: str
    condition: str
    group: str
    ires_type: str
    organism: str
    activity_label: int | None
    activity_status: str
    te_ires: float | None
    te_cdi: float | None
    fold_difference_from_wt: float | None
    position: float | None
    is_control: bool
    notes: str
    pmid: str

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    return str(value).strip()


def _optional_float(value: Any) -> float | None:
    if value is None or _clean_text(value) == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def interpret_activity(value: Any) -> tuple[int | None, str]:
    """Map only literal validated active/inactive calls to binary labels.

    Qualified calls such as ``active (not tested)`` and ``active (does not
    validate)`` are retained as ambiguous rather than silently promoted to
    positives.
    """
    raw = _clean_text(value).lower()
    if raw == "active":
        return 1, "active"
    if raw == "inactive":
        return 0, "inactive"
    if raw:
        return None, f"ambiguous:{raw}"
    return None, "unlabeled"


def _sequence_hash(sequence: str) -> str:
    return hashlib.sha256(sequence.encode("ascii", errors="ignore")).hexdigest()


def _row_value(row: dict[str, Any], *names: str) -> Any:
    for name in names:
        if name in row:
            return row[name]
    return None


def _make_measurement(
    *,
    source_table: str,
    source_row: int,
    row: dict[str, Any],
    condition: str = "baseline",
    te_ires: Any = None,
    te_cdi: Any = None,
    activity: Any = None,
) -> TrapprMeasurement | None:
    sequence = normalize_rna(
        _clean_text(_row_value(row, "Sequence", "sequence"))
    )
    if not sequence:
        return None
    construct = _clean_text(_row_value(row, "Construct name", "Construct Name"))
    group = _clean_text(_row_value(row, "group", "Category"))
    ires_type = _clean_text(_row_value(row, "Type"))
    organism = _clean_text(_row_value(row, "Isolated from"))
    notes = _clean_text(_row_value(row, "Notes"))
    pmid = _clean_text(_row_value(row, "PMID (IRESbase)", "PubmedID"))
    label, status = interpret_activity(
        activity if activity is not None else _row_value(row, "IRES activity", "IRES Activity")
    )
    control_text = " ".join((construct, group, ires_type)).lower()
    return TrapprMeasurement(
        record_id=f"trappr_{source_table}_r{source_row:05d}_{condition}",
        source_table=source_table,
        source_row=source_row,
        construct_name=construct,
        sequence=sequence,
        sequence_sha256=_sequence_hash(sequence),
        length=len(sequence),
        alphabet_valid=set(sequence) <= set("ACGU"),
        assay_context=ASSAY_CONTEXT,
        condition=condition,
        group=group,
        ires_type=ires_type,
        organism=organism,
        activity_label=label,
        activity_status=status,
        te_ires=_optional_float(
            te_ires if te_ires is not None else _row_value(row, "TE - IRES")
        ),
        te_cdi=_optional_float(te_cdi if te_cdi is not None else _row_value(row, "TE - CDI")),
        fold_difference_from_wt=_optional_float(_row_value(row, "Fold-Difference-From-WT")),
        position=_optional_float(_row_value(row, "HCV-position")),
        is_control="control" in control_text or construct.lower().startswith("g-cap"),
        notes=notes,
        pmid=pmid,
    )


def load_trappr_workbooks(root: Path) -> list[TrapprMeasurement]:
    """Load sequence-bearing supplementary Tables S1--S7.

    Table S7 contains paired control/stress measurements, which are emitted as
    two records with identical sequence provenance and distinct conditions.
    Tables S8 and S9 contain construct/primer metadata and construct-level
    validation summaries without a sequence column, so they are audited but not
    coerced into sequence measurements here.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - optional dependency guard
        raise RuntimeError("openpyxl is required to normalize IRES-TrAPPr XLSX files") from exc

    records: list[TrapprMeasurement] = []
    for table_number in range(1, 8):
        path = root / f"media-{table_number}.xlsx"
        if not path.exists():
            raise FileNotFoundError(path)
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook[workbook.sheetnames[0]]
        rows = worksheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        source_table = f"S{table_number}"
        for source_row, values in enumerate(rows, start=2):
            row = dict(zip(headers, values))
            if table_number == 7:
                for condition, te_i, te_c in (
                    ("control", row.get("Control TE - IRES"), row.get("Control TE- CDI")),
                    ("stress", row.get("Stress TE - IRES"), row.get("Stress TE - CDI")),
                ):
                    record = _make_measurement(
                        source_table=source_table,
                        source_row=source_row,
                        row=row,
                        condition=condition,
                        te_ires=te_i,
                        te_cdi=te_c,
                    )
                    if record is not None:
                        records.append(record)
            else:
                record = _make_measurement(
                    source_table=source_table,
                    source_row=source_row,
                    row=row,
                )
                if record is not None:
                    records.append(record)
        workbook.close()
    return records


def build_sequence_view(records: Iterable[TrapprMeasurement]) -> list[dict[str, Any]]:
    grouped: dict[str, list[TrapprMeasurement]] = defaultdict(list)
    for record in records:
        grouped[record.sequence_sha256].append(record)

    output: list[dict[str, Any]] = []
    for sequence_hash, rows in sorted(grouped.items()):
        labels = [row.activity_label for row in rows if row.activity_label is not None]
        label_set = set(labels)
        te_all = sorted({row.te_ires for row in rows if row.te_ires is not None})
        te_baseline = sorted(
            {
                row.te_ires
                for row in rows
                if row.te_ires is not None and row.condition in {"baseline", "control"}
            }
        )
        te_stress = sorted(
            {row.te_ires for row in rows if row.te_ires is not None and row.condition == "stress"}
        )
        cdi_all = sorted({row.te_cdi for row in rows if row.te_cdi is not None})
        fold_changes = sorted(
            {
                row.fold_difference_from_wt
                for row in rows
                if row.fold_difference_from_wt is not None
            }
        )
        consensus_label: int | None = next(iter(label_set)) if len(label_set) == 1 else None
        label_status = (
            "conflict"
            if len(label_set) > 1
            else "active"
            if consensus_label == 1
            else "inactive"
            if consensus_label == 0
            else "unlabeled_or_ambiguous"
        )
        output.append(
            {
                "sequence_id": f"trappr_seq_{sequence_hash[:16]}",
                "sequence": rows[0].sequence,
                "sequence_sha256": sequence_hash,
                "length": rows[0].length,
                "alphabet_valid": rows[0].alphabet_valid,
                "assay_context": ASSAY_CONTEXT,
                "n_measurements": len(rows),
                "n_active": labels.count(1),
                "n_inactive": labels.count(0),
                "n_ambiguous": sum(row.activity_status.startswith("ambiguous:") for row in rows),
                "n_unlabeled": sum(row.activity_status == "unlabeled" for row in rows),
                "consensus_label": consensus_label,
                "consensus_status": label_status,
                "te_ires_mean_unique": fmean(te_all) if te_all else None,
                "te_ires_max": max(te_all) if te_all else None,
                "te_ires_baseline_mean_unique": fmean(te_baseline) if te_baseline else None,
                "te_ires_stress_mean_unique": fmean(te_stress) if te_stress else None,
                "te_cdi_mean_unique": fmean(cdi_all) if cdi_all else None,
                "fold_difference_from_wt_mean_unique": (
                    fmean(fold_changes) if fold_changes else None
                ),
                "is_control": any(row.is_control for row in rows),
                "construct_names": "|".join(sorted({row.construct_name for row in rows if row.construct_name})),
                "source_tables": "|".join(sorted({row.source_table for row in rows})),
                "groups": "|".join(sorted({row.group for row in rows if row.group})),
                "ires_types": "|".join(sorted({row.ires_type for row in rows if row.ires_type})),
                "organisms": "|".join(sorted({row.organism for row in rows if row.organism})),
            }
        )
    return output


def audit_trappr(
    records: list[TrapprMeasurement], sequence_rows: list[dict[str, Any]]
) -> dict[str, Any]:
    statuses = Counter(record.activity_status for record in records)
    consensus = Counter(row["consensus_status"] for row in sequence_rows)
    return {
        "schema_version": 1,
        "assay_context": ASSAY_CONTEXT,
        "n_measurements": len(records),
        "n_unique_sequences": len(sequence_rows),
        "n_controls": sum(record.is_control for record in records),
        "n_invalid_alphabet_measurements": sum(not record.alphabet_valid for record in records),
        "measurement_count_by_table": dict(sorted(Counter(r.source_table for r in records).items())),
        "measurement_count_by_condition": dict(sorted(Counter(r.condition for r in records).items())),
        "measurement_activity_status": dict(sorted(statuses.items())),
        "sequence_consensus_status": dict(sorted(consensus.items())),
        "sequence_length": {
            "min": min(row["length"] for row in sequence_rows),
            "max": max(row["length"] for row in sequence_rows),
            "mean": sum(row["length"] for row in sequence_rows) / len(sequence_rows),
        },
        "excluded_from_sequence_measurements": {
            "S8": "gene fragments and primers; no sequence-level activity target",
            "S9": "construct-level luciferase validation; join by construct name in a later validation view",
        },
    }


def _write_dict_rows(rows: list[dict[str, Any]], path: Path, fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_normalized_trappr(root: Path, output_dir: Path) -> dict[str, Any]:
    records = load_trappr_workbooks(root)
    sequence_rows = build_sequence_view(records)
    audit = audit_trappr(records, sequence_rows)
    _write_dict_rows(
        [record.to_dict() for record in records],
        output_dir / "measurements.csv",
        list(MEASUREMENT_FIELDS),
    )
    _write_dict_rows(
        sequence_rows,
        output_dir / "sequences.csv",
        list(sequence_rows[0]),
    )
    (output_dir / "audit.json").write_text(
        json.dumps(audit, indent=2, ensure_ascii=False, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return audit
